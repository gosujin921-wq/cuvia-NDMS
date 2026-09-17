# -*- coding: utf-8 -*-
"""CUVIA_NDMS_기능및정보구조도.md 를 읽어 같은 이름의 xlsx 로 옮긴다.

사용: python3 build_ia_xlsx.py <IA.md> <출력.xlsx>

내용은 md 를 파싱해서 만든다. 문구를 손으로 다시 적지 않으므로 md 를 고치고 이 스크립트를
다시 돌리면 엑셀이 따라온다. 시트 10개 구성은 IA md §2 가 정한 것이다.
값이 갈리면 md 가 이긴다.
"""
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]

# ───────────────────────── md 파싱

def strip_md(t):
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = re.sub(r"\*\*([^*]*)\*\*", r"\1", t)
    t = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", t)
    return t.strip()


def parse_blocks(lines):
    """줄 목록을 블록으로. 종류: table(rows) · list(items) · para(text) · quote(text) · code(text)"""
    blocks, i, n = [], 0, len(lines)
    while i < n:
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("```"):
            j = i + 1
            buf = []
            while j < n and not lines[j].startswith("```"):
                buf.append(lines[j])
                j += 1
            blocks.append(("code", "\n".join(buf)))
            i = j + 1
            continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                cells = [strip_md(c) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if ln.startswith("> "):
            buf = []
            while i < n and lines[i].startswith(">"):
                buf.append(lines[i][1:].strip())
                i += 1
            blocks.append(("quote", strip_md(" ".join(b for b in buf if b))))
            continue
        m = re.match(r"^(\s*)(?:[-*]|\d+\.)\s+(.*)", ln)
        if m:
            items = []
            while i < n:
                m2 = re.match(r"^(\s*)(?:[-*]|\d+\.)\s+(.*)", lines[i])
                if m2:
                    items.append(m2.group(2))
                    i += 1
                elif lines[i].startswith("  ") and lines[i].strip() and items:
                    items[-1] += " " + lines[i].strip()
                    i += 1
                else:
                    break
            blocks.append(("list", [strip_md(x) for x in items]))
            continue
        buf = []
        while i < n and lines[i].strip() and not re.match(r"^(#|\||>|```|\s*[-*]\s|\s*\d+\.\s)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("para", strip_md(" ".join(buf))))
    return blocks


def parse_md(path):
    """{절번호: {"title":, "sub": [(h3, h4, blocks)]}}. h3/h4 가 없으면 ''."""
    text = open(path, encoding="utf8").read()
    lines = text.split("\n")
    secs, cur, h3, h4, buf = {}, None, "", "", []

    def flush():
        if cur is not None:
            secs[cur]["sub"].append((h3, h4, parse_blocks(buf)))

    for ln in lines:
        m = re.match(r"^## (\d+)\. (.*)", ln)
        if m:
            flush()
            cur, h3, h4, buf = int(m.group(1)), "", "", []
            secs[cur] = {"title": m.group(2).strip(), "sub": []}
            continue
        m = re.match(r"^### (.*)", ln)
        if m and cur is not None:
            flush()
            h3, h4, buf = m.group(1).strip(), "", []
            continue
        m = re.match(r"^#### (.*)", ln)
        if m and cur is not None:
            flush()
            h4, buf = m.group(1).strip(), []
            continue
        if cur is not None:
            buf.append(ln)
    flush()
    head = {}
    for ln in lines[:12]:
        m = re.match(r"^> (상태|기준일|역할): (.*)", ln)
        if m:
            head[m.group(1)] = m.group(2).strip()
    return secs, head


def blocks_of(secs, num, h3="", h4=None):
    out = []
    for a, b, blocks in secs[num]["sub"]:
        if a == h3 and (h4 is None or b == h4):
            out.extend(blocks)
    return out


def first_table(blocks):
    for k, v in blocks:
        if k == "table":
            return v
    return []


# ───────────────────────── xlsx

HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
HEAD_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
wb = Workbook()
wb.remove(wb.active)


def sheet(title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, WRAP, BORDER
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


secs, head = parse_md(SRC)
stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

# 0. 개요 (§1 · §2 · §3 · §21)
rows = [
    ("문서", "원본", f"{SRC.split('/')[-1]} · 이 엑셀은 md 를 표로 옮긴 동기화본 ({stamp} 생성). 값이 갈리면 md 가 이긴다"),
    ("문서", "상태", head.get("상태", "")),
    ("문서", "기준일", head.get("기준일", "")),
    ("문서", "역할", head.get("역할", "")),
]
for k, v in blocks_of(secs, 1):
    if k == "quote":
        rows.append(("§1 목적", "정의", v))
    elif k == "para":
        rows.append(("§1 목적", "", v))
for k, v in blocks_of(secs, 2):
    if k == "para":
        rows.append(("§2 엑셀 역할", "", v))
    elif k == "list":
        rows.extend(("§2 엑셀 역할", "", x) for x in v)
for k, v in blocks_of(secs, 3):
    if k == "list":
        for x in v:
            m = re.match(r"^([^:]+):\s*(.*)", x)
            rows.append(("§3 원칙", m.group(1), m.group(2)) if m else ("§3 원칙", "", x))
for k, v in blocks_of(secs, 21):
    if k == "list":
        rows.extend(("§21 완료 기준", "", x) for x in v)
sheet("개요", ["구분", "항목", "내용"], rows, [22, 20, 100])

# 1. 화면 (§5 표 + §4 트리 하위 구성 + §5.1 진입 조건)
tree = ""
for k, v in blocks_of(secs, 4):
    if k == "code":
        tree = v
children = {}
cur = None
for ln in tree.split("\n"):
    m = re.match(r"^[│ ]*[├└]── (\d\d) (.*)", ln)
    if m:
        cur = m.group(2).strip()
        children[cur] = []
        continue
    m = re.match(r"^[│ ]*[├└]── (.*)", ln)
    if m and cur and ln.startswith("│"):
        children[cur].append(m.group(1).strip())
    elif m:
        cur = None
entry = {r[0]: r for r in first_table(blocks_of(secs, 5, "5.1 진입 조건과 다음 행동"))[1:]}
route_table = first_table(blocks_of(secs, 5, "5.2 Phase 2 라우트 계약"))
routes = {r[0]: r[1] for r in route_table[1:]}
rows = []
for r in first_table(blocks_of(secs, 5))[1:]:
    name = r[1]
    ent = next((v for k, v in entry.items() if name in k), ("", "", ""))
    rows.append((r[0], name, r[2], routes.get(r[0], ""), r[3], r[4], r[5], " / ".join(children.get(name, [])), ent[1], ent[2]))
for k, v in blocks_of(secs, 5):
    if k == "para":
        rows.append(("", "", "", "", v, "", "", "", "", ""))
sheet("화면", ["ID", "업무 공간", "제공 형태", "기준 라우트", "담당자의 핵심 질문", "대표 산출물", "데모 위치", "하위 구성 (§4)", "주 진입 조건 (§5.1)", "다음 행동 (§5.1)"],
      rows, [9, 16, 24, 34, 40, 24, 12, 40, 36, 36])

# 2. 정보구조도 (§6~§12 본문 전부)
R = []
twin_table = []
for num in range(6, 13):
    title = secs[num]["title"]
    m = re.match(r"^(IA-\S+)\s+(.*)", title)
    iid, space = (m.group(1), m.group(2)) if m else ("", title)
    for h3, h4, blocks in secs[num]["sub"]:
        for k, v in blocks:
            if k == "table":
                hdr, body = v[0], v[1:]
                if h3 == "유형별 교체 영역":
                    twin_table = v
                    continue
                for r in body:
                    obj = r[2] if len(r) > 2 else ""
                    R.append((iid, space, h3, r[0] if h4 == "" else h4 + " · " + r[0], " | ".join(r[1:2]) if len(r) > 1 else "", obj))
            elif k == "list":
                for x in v:
                    R.append((iid, space, h3, h4, x, ""))
            elif k in ("para", "quote"):
                R.append((iid, space, h3, h4, v, ""))
sheet("정보구조도", ["ID", "화면", "구분", "영역", "내용", "원천 객체"], R, [9, 14, 16, 18, 70, 24])

# 3. 트윈유형 (§8 유형별 교체 영역 표 + 뒤따르는 설명)
rows = [tuple(r) for r in twin_table[1:]]
notes = [v for k, v in blocks_of(secs, 8, "유형별 교체 영역") if k == "para"]
for t in notes:
    rows.append((t, "", "", "", "", ""))
sheet("트윈유형", twin_table[0] if twin_table else ["비교 기준안"], rows, [22, 26, 16, 26, 12, 30])

# 4. 데모동선 (§13)
rows = []
for r in first_table(blocks_of(secs, 13))[1:]:
    m = re.match(r"^(D\d)\s+(.*)", r[0])
    rows.append(((m.group(1), m.group(2)) if m else (r[0], "")) + tuple(r[1:]))
for k, v in blocks_of(secs, 13):
    if k == "para":
        rows.append(("", "", "", v, "", ""))
sheet("데모동선", ["시나리오", "단계", "화면", "화면에서 답할 질문", "주요 행위", "기준 라우트"], rows, [10, 14, 26, 40, 40, 46])

# 5. 사건맥락 (§14)
rows = [tuple(r) for r in first_table(blocks_of(secs, 14))[1:]]
for k, v in blocks_of(secs, 14):
    if k == "para":
        rows.append(("", v))
sheet("사건맥락", ["맥락", "유지 이유"], rows, [24, 70])

# 6. 정보객체 (§15)
rows = [tuple(r) for r in first_table(blocks_of(secs, 15))[1:]]
for k, v in blocks_of(secs, 15):
    if k == "para":
        rows.append(("", "", v))
sheet("정보객체", ["객체", "사용하는 공간", "핵심 정보"], rows, [22, 18, 70])

# 7. 공통상태 (§16)
rows = [tuple(r) for r in first_table(blocks_of(secs, 16))[1:]]
sheet("공통상태", ["상태", "화면 처리"], rows, [18, 70])

# 8. Phase1판정 (§17)
rows = [tuple(r) for r in first_table(blocks_of(secs, 17))[1:]]
sheet("Phase1판정", ["Phase 1 요소", "Phase 2 판정", "이동 대상"], rows, [30, 14, 44])

# 9. 데모범위 (§18 · §19 · §20)
rows = []
for num in (18, 19, 20):
    sec = secs[num]["title"]
    for h3, h4, blocks in secs[num]["sub"]:
        label = f"§{num} {sec}" + (f" · {h3}" if h3 else "")
        for k, v in blocks:
            if k == "list":
                rows.extend((label, x) for x in v)
            elif k in ("para", "quote"):
                rows.append((label, v))
sheet("데모범위", ["구분", "내용"], rows, [34, 80])

wb.save(OUT)
print("저장:", OUT)
for ws in wb.worksheets:
    print(f"  {ws.title}: {ws.max_row - 1}행")
